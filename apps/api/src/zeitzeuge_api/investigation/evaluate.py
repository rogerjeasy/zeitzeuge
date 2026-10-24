"""Score exported workbooks against the team's labelled workbook (ground truth).

    python -m zeitzeuge_api.investigation.evaluate GOLD.xlsx EXPORT.xlsx [EXPORT2.xlsx ...]

GOLD is a copy of the Ergebnisse sheet in which the team coloured every correct
value green (fill FF92D050) and corrected the values that were wrong. Only green
cells are truth; an empty or uncoloured cell means «not labelled» and is not
scored. Each EXPORT is scored per ranked field:

    correct   the export's value names the labelled value
    wrong     the export filled the field with something else
    missed    the export left the field empty although a label exists

Matching follows the pipeline's own name rules (identification-upgrade note §12):
case, accents and «St./Sankt» folded; generic building words dropped («Hotel
Restaurant Seehof» names the same house as «Haus du Lac Seehof»); a street
without a number matches the same street with one; small spelling slips in a
label are tolerated (similarity ≥ 0.85).

Labels are checked too: a «Gemeinde» label that is not one of the canton's 30
political municipalities (Stoos, Goldau, …) is reported as a label issue and
excluded from the Gemeinde score instead of being counted against the export.
"""

from __future__ import annotations

import argparse
import re
import sys
from collections.abc import Iterable
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import IO

from openpyxl import load_workbook

from zeitzeuge_api.investigation.names import distinctive, feature_key
from zeitzeuge_api.investigation.tools.text import fold, tokens

GREEN = "FF92D050"
FIELDS: tuple[tuple[str, str], ...] = (
    ("Ortschaft", "Ortschaft"),
    ("Gemeinde", "Gemeinde"),
    ("Gebaeude", "Gebäude"),
    ("See_Fluss_Berg", "See / Fluss / Berg"),
    ("Strasse", "Strasse"),
)
MUNICIPALITIES = frozenset(
    fold(n)
    for n in (
        "Alpthal", "Altendorf", "Arth", "Einsiedeln", "Feusisberg", "Freienbach", "Galgenen",
        "Gersau", "Illgau", "Ingenbohl", "Innerthal", "Küssnacht", "Lachen", "Lauerz",
        "Morschach", "Muotathal", "Oberiberg", "Reichenburg", "Riemenstalden", "Rothenthurm",
        "Sattel", "Schübelbach", "Schwyz", "Steinen", "Steinerberg", "Tuggen", "Unteriberg",
        "Vorderthal", "Wangen", "Wollerau",
    )
)  # fmt: skip
_PLACE_NOISE = {"am", "sz", "bei", "im", "rigi"}
_ABBREVIATIONS = ((r"\bg\b\.?", "gasse"), (r"\bstr\b\.?", "strasse"), (r"\bst\b\.?", "sankt"))
SIMILAR = 0.85


@dataclass(slots=True)
class FieldScore:
    labelled: int = 0
    correct: int = 0
    wrong: int = 0
    missed: int = 0
    wrong_cases: list[str] = field(default_factory=list)
    missed_cases: list[str] = field(default_factory=list)

    @property
    def precision(self) -> float | None:
        filled = self.correct + self.wrong
        return self.correct / filled if filled else None

    @property
    def recall(self) -> float | None:
        return self.correct / self.labelled if self.labelled else None


@dataclass(slots=True)
class Report:
    name: str
    fields: dict[str, FieldScore]
    bins: dict[str, list[int]]
    """«0.90-1.00» → [fields, correct]."""
    label_issues: list[str]


def _expand(value: str) -> str:
    text = fold(value)
    for pattern, full in _ABBREVIATIONS:
        text = re.sub(pattern, full, text)
    return text


def _similar(a: str, b: str) -> bool:
    return SequenceMatcher(None, a, b).ratio() >= SIMILAR


def same(column: str, predicted: str, label: str) -> bool:
    """Does the export's value name the labelled thing?"""
    p, gold = _expand(predicted), _expand(label)
    if not p or not gold:
        return False
    if p == gold or _similar(p, gold):
        return True
    if column == "See_Fluss_Berg":
        return feature_key(predicted) == feature_key(label) or _similar(
            feature_key(predicted), feature_key(label)
        )
    if column == "Gebaeude":
        core_p, core_g = distinctive(p), distinctive(gold)
        if not core_p or not core_g:
            return False
        tp, tg = set(core_p.split()), set(core_g.split())
        return (tp <= set(tokens(p)) and (tp <= tg or tg <= tp)) or _similar(core_p, core_g)
    if column == "Strasse":
        name_p = " ".join(t for t in p.split() if not t[0].isdigit())
        name_g = " ".join(t for t in gold.split() if not t[0].isdigit())
        num_p = [t for t in p.split() if t[0].isdigit()]
        num_g = [t for t in gold.split() if t[0].isdigit()]
        street_ok = name_p == name_g or _similar(name_p, name_g) or name_p.startswith(name_g)
        return street_ok and (not num_g or num_p[:1] == num_g[:1])
    tp = set(tokens(p)) - _PLACE_NOISE
    tg = set(tokens(gold)) - _PLACE_NOISE
    return bool(tp) and bool(tg) and (tp <= tg or tg <= tp)


def _rgb(cell: object) -> str:
    fill = getattr(cell, "fill", None)
    if fill is None or not fill.fill_type or fill.fgColor.type != "rgb":
        return ""
    return str(fill.fgColor.rgb)


MAX_HEADER_COLUMNS = 128
"""The export has a few dozen columns; anything wider is not one of ours."""


def read_labels(path: Path | IO[bytes], max_rows: int | None = None) -> dict[str, dict[str, str]]:
    """The green cells of the first sheet, per photo (`Dateiname`).

    Streams the rows (read-only mode keeps the fills) and reads only the
    columns it needs, so an uploaded workbook is never materialised whole.
    openpyxl parses with defusedxml (a dependency), so entity-expansion tricks
    are refused.
    """
    book = load_workbook(path, read_only=True)
    try:
        sheet = book.worksheets[0]
        rows = sheet.iter_rows(min_row=1, max_col=MAX_HEADER_COLUMNS)
        header = [c.value for c in next(rows, ())]
        missing = [c for c in ("Dateiname", *(column for column, _ in FIELDS)) if c not in header]
        if missing:
            raise ValueError(f"Not a ZEITZEUGE results sheet: missing {', '.join(missing)}.")
        name_at = header.index("Dateiname")
        wanted = {column: header.index(column) for column, _ in FIELDS}
        width = max(name_at, *wanted.values()) + 1
        labels: dict[str, dict[str, str]] = {}
        for count, row in enumerate(sheet.iter_rows(min_row=2, max_col=width), start=1):
            if max_rows is not None and count > max_rows:
                raise ValueError(f"The sheet has more than {max_rows} rows.")
            if len(row) <= name_at or not row[name_at].value:
                continue
            labels[str(row[name_at].value)] = {
                column: str(row[at].value).strip()
                for column, at in wanted.items()
                if at < len(row) and row[at].value and _rgb(row[at]) == GREEN
            }
        return labels
    finally:
        book.close()


def read_export(path: Path) -> dict[str, dict[str, tuple[str, float | None]]]:
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book["Ergebnisse"] if "Ergebnisse" in book.sheetnames else book.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    out: dict[str, dict[str, tuple[str, float | None]]] = {}
    for row in rows:
        if not row or not row[0]:
            continue
        values: dict[str, tuple[str, float | None]] = {}
        for column, _ in FIELDS:
            value = row[header.index(column)]
            confidence = row[header.index(f"{column}_Konfidenz")]
            if value:
                values[column] = (str(value), float(confidence) if confidence else None)
        out[str(row[0])] = values
    return out


def label_issues(labels: dict[str, dict[str, str]]) -> dict[str, set[str]]:
    """Labels that cannot be right by definition, per photo: which columns to skip."""
    issues: dict[str, set[str]] = {}
    for photo, values in labels.items():
        municipality = values.get("Gemeinde")
        if municipality and fold(municipality) not in MUNICIPALITIES:
            issues.setdefault(photo, set()).add("Gemeinde")
    return issues


def _bin(confidence: float | None) -> str:
    if confidence is None:
        return "ohne Zahl"
    low = min(int(confidence * 10), 9) / 10
    return f"{max(low, 0.5):.2f}-{min(low + 0.1, 1.0):.2f}" if low >= 0.5 else "unter 0.50"


def score(
    name: str,
    labels: dict[str, dict[str, str]],
    export: dict[str, dict[str, tuple[str, float | None]]],
) -> Report:
    skip = label_issues(labels)
    fields = {column: FieldScore() for column, _ in FIELDS}
    bins: dict[str, list[int]] = {}
    for photo, gold in labels.items():
        predicted = export.get(photo, {})
        for column, truth in gold.items():
            if column in skip.get(photo, set()):
                continue
            result = fields[column]
            result.labelled += 1
            short = photo.replace("SG.CIV.17.", "").replace(".jpg", "")
            if column not in predicted:
                result.missed += 1
                result.missed_cases.append(f"{short} {column}: «{truth}»")
                continue
            value, confidence = predicted[column]
            ok = same(column, value, truth)
            if ok:
                result.correct += 1
            else:
                result.wrong += 1
                result.wrong_cases.append(f"{short} {column}: «{value}» statt «{truth}»")
            slot = bins.setdefault(_bin(confidence), [0, 0])
            slot[0] += 1
            slot[1] += int(ok)
    issues = [
        f"{photo.replace('.jpg', '')}: Gemeinde «{labels[photo]['Gemeinde']}» ist keine "
        "politische Gemeinde des Kantons"
        for photo, columns in skip.items()
        if "Gemeinde" in columns
    ]
    return Report(name=name, fields=fields, bins=bins, label_issues=issues)


def _pct(value: float | None) -> str:
    return "—" if value is None else f"{value:.0%}"


def render(reports: Iterable[Report]) -> str:
    reports = list(reports)
    lines: list[str] = []
    for report in reports:
        totals = FieldScore()
        lines.append(f"\n## {report.name}\n")
        lines.append("| Feld | gelabelt | richtig | falsch | leer | Präzision | Abdeckung |")
        lines.append("|---|---|---|---|---|---|---|")
        for column, label in FIELDS:
            s = report.fields[column]
            for attr in ("labelled", "correct", "wrong", "missed"):
                setattr(totals, attr, getattr(totals, attr) + getattr(s, attr))
            lines.append(
                f"| {label} | {s.labelled} | {s.correct} | {s.wrong} | {s.missed} | "
                f"{_pct(s.precision)} | {_pct(s.recall)} |"
            )
        lines.append(
            f"| **Alle** | {totals.labelled} | {totals.correct} | {totals.wrong} | "
            f"{totals.missed} | {_pct(totals.precision)} | {_pct(totals.recall)} |"
        )
        lines.append("\n| Konfidenz | Felder | richtig | Trefferquote |\n|---|---|---|---|")
        for key in sorted(report.bins, reverse=True):
            n, ok = report.bins[key]
            lines.append(f"| {key} | {n} | {ok} | {_pct(ok / n if n else None)} |")
        wrong = [c for column, _ in FIELDS for c in report.fields[column].wrong_cases]
        if wrong:
            lines.append("\nFalsch: " + "; ".join(wrong))
    if reports and reports[0].label_issues:
        lines.append("\nLabel-Hinweise (nicht gewertet): " + "; ".join(reports[0].label_issues))
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("gold", type=Path)
    parser.add_argument("exports", type=Path, nargs="+")
    parser.add_argument(
        "--only",
        type=Path,
        help="score only the photos present in this export (compare runs on the same subset)",
    )
    args = parser.parse_args(argv)
    labels = read_labels(args.gold)
    if args.only:
        present = set(read_export(args.only))
        labels = {photo: values for photo, values in labels.items() if photo in present}
    reports = [score(path.name, labels, read_export(path)) for path in args.exports]
    sys.stdout.write(
        f"# Auswertung gegen {args.gold.name}: {len(labels)} gelabelte Fotos\n"
        + render(reports)
        + "\n"
    )


if __name__ == "__main__":
    main()
