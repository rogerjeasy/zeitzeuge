"""The Excel workbook: one row per photograph (data-model §8.12, design note §4).

This is the brief's mandatory deliverable, and the person who reads it wrote
the brief. Three properties matter more than anything else here.

* **One row per delivered file, always.** A photo whose run has not finished
  still gets its row (file name, SHA-256, Prio, `QC_Status = offen`), so the
  archive can check the count against what it delivered. The workbook has to
  exist before the agents do.
* **Every number explains itself.** Each ranked field carries its own
  confidence and its own source line, every cited item is on «Belege» with a
  link, and «Methode» prints the rule with the constants read from
  `investigation/confidence.py`, so the sheet and the code cannot disagree.
* **Empty means empty.** An abstained field has an empty value *and* an empty
  confidence cell, never 0; camera columns stay empty until a pose is
  resolved. Nothing is filled in to look complete.

Columns are data: the ranked fields' names and order come from
`field_definitions` (`excel_column`, `priority_rank`), so a later import format
is a mapping edit on «Spalten», not code.

Everything a model or a website wrote is untrusted text. It is written as a
string cell, never a formula, with characters Excel cannot store removed, and
only http(s) URLs become hyperlinks.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import Collection, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from zeitzeuge_api.db.models.catalogue import FieldDefinition, RecordField, RecordFieldEvidence
from zeitzeuge_api.db.models.enums import FieldDecisionKind, SourceClass
from zeitzeuge_api.db.models.investigation import EvidenceItem
from zeitzeuge_api.investigation.confidence import (
    ACCEPT_AT,
    MODEL_STRENGTH_CAP,
    OPEN_WEB_PAGE,
    REFERENCE_TITLE_MATCH,
    REVIEW_AT,
    SOURCE_CLASS_DE,
    WEIGHTS,
    Support,
    score,
    strength_for_match,
)
from zeitzeuge_api.investigation.tools.base import MatchKind
from zeitzeuge_api.services.qc import (
    QC_OPEN,
    RANKED_FIELD_KEYS,
    AssetState,
    BatchState,
    load_batch,
)

STATUS_ACCEPT = "übernahmefähig"
STATUS_REVIEW = "Prüfbedarf"
STATUS_NONE = "unbestimmt"
"""`unbestimmt`, not `offen`, so the rule's verdict cannot be confused with
`QC_Status` (design note §4)."""

SHEETS = ("Ergebnisse", "Belege", "Methode", "Spalten")

_FALLBACK_COLUMNS: Mapping[str, tuple[str, str]] = {
    "locality": ("Ortschaft", "Ortschaft"),
    "municipality": ("Gemeinde", "Politische Gemeinde"),
    "building": ("Gebaeude", "Gebäude"),
    "landscape_feature": ("See_Fluss_Berg", "See / Fluss / Berg"),
    "street": ("Strasse", "Strasse"),
}
"""Only used if `field_definitions` lacks a ranked row (an unseeded database):
the deliverable must still have its columns rather than fail."""

_LEADING = ("Dateiname", "Datei_SHA256", "Signatur", "Prio")
_TRAILING = (
    "Motiv_LV95_E",
    "Motiv_LV95_N",
    "Aufnahmestandort_LV95_E",
    "Aufnahmestandort_LV95_N",
    "Aufnahmerichtung_Grad",
    "Aufnahmedatum_EXIF",
    "Ausserhalb_SZ",
    "Gesamtkonfidenz",
    "Status",
    "Pruefbedarf",
    "Unbestimmt",
    "QC_Status",
    "QC_Von",
    "Bemerkung",
    "Run_ID",
    "Erzeugt_am",
)

_BELEGE_HEADER = (
    "Dateiname",
    "Feld",
    "Quellenart",
    "Quelle",
    "URL",
    "Beobachtung",
    "Staerke",
    "Agent",
    "Modell",
)

_MINUS = "\u2212"
_TIMES = "\u00d7"
"""Typographic minus and multiplication signs for the formula on «Methode»."""

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="E7E6E6")
_TITLE_FONT = Font(bold=True, size=14)
_SECTION_FONT = Font(bold=True, size=12)
_WRAP = Alignment(wrap_text=True, vertical="top")
_TOP = Alignment(vertical="top")
_LINK_FONT = Font(color="0563C1", underline="single")

# Light fills with dark text of the same hue: readable on screen, in print and
# for most colour-vision deficiencies, because the word carries the meaning.
_STATUS_STYLE: Mapping[str, tuple[PatternFill, Font]] = {
    STATUS_ACCEPT: (PatternFill("solid", fgColor="E2F0D9"), Font(color="375623")),
    STATUS_REVIEW: (PatternFill("solid", fgColor="FFF2CC"), Font(color="7F6000")),
    STATUS_NONE: (PatternFill("solid", fgColor="EDEDED"), Font(color="404040")),
}

_ABSTAIN_DE: Mapping[str, str] = {
    "insufficient_evidence": "zu wenige oder zu schwache Belege",
    "conflicting_evidence": "widersprüchliche Belege oder widerlegt",
    "out_of_scope": "im Bild nicht anwendbar",
    "sensitivity_hold": "aus Schutzgründen zurückgehalten",
}

_SOURCE_CLASS_SCOPE: Mapping[SourceClass, str] = {
    SourceClass.EXIF: "Aufnahmedaten in der Bilddatei (Datum, Kamera, Brennweite).",
    SourceClass.VISUAL: "Was ein Bildmodell im Foto erkennt: Gebäudetyp, Landschaft, Wahrzeichen.",
    SourceClass.TEXT_IN_IMAGE: "Lesbare Schrift im Bild: Schilder, Strassentafeln, Beschriftungen.",
    SourceClass.GEODATA_LAYER: (
        "Amtliche und offene Geodaten: swisstopo (Ortschaftenverzeichnis, swissBOUNDARIES3D, "
        "swissNAMES3D), Gebäudeadressverzeichnis/GWR, OpenStreetMap."
    ),
    SourceClass.REFERENCE_IMAGE: (
        "Vergleichsbilder desselben Motivs: Katalog des Staatsarchivs (ohne die Serie "
        "SG.CIV.17, damit die Lösung nicht aus dem eigenen Bestand stammt), Wikimedia Commons."
    ),
    SourceClass.OPEN_WEB: "Webseiten, Wikipedia, Wikidata.",
    SourceClass.TERRAIN_GEOMETRY: "Abgleich der Horizontlinie mit dem Geländemodell (noch nicht "
    "in Betrieb).",
    SourceClass.STREET_LEVEL: "Abgleich mit Strassenansichten (noch nicht in Betrieb).",
    SourceClass.ARCHIVAL: "Archivkataloge und Findmittel.",
    SourceClass.HUMAN: "Angabe oder Korrektur durch eine Archivarin oder einen Archivar.",
}


@dataclass(frozen=True, slots=True)
class _Ranked:
    key: str
    column: str
    label: str
    description: str | None


async def build_workbook(
    session: AsyncSession,
    batch_id: uuid.UUID,
    *,
    generated_at: datetime | None = None,
    asset_ids: Collection[uuid.UUID] | None = None,
) -> bytes:
    """The `.xlsx` for one batch, read under the session's tenant.

    Raises `BatchNotFoundError` if the batch is not visible.
    """
    stamp = _iso(generated_at or datetime.now(UTC))
    state = await load_batch(session, batch_id, asset_ids=asset_ids)
    ranked = _ranked_fields(state.definitions)
    header = _ergebnisse_header(ranked)
    evidence = await _cited_evidence(session, state)

    workbook = Workbook()
    ergebnisse = workbook.active
    assert ergebnisse is not None
    ergebnisse.title = SHEETS[0]
    _write_ergebnisse(ergebnisse, header, [_row(a, ranked, stamp) for a in state.assets])
    _write_belege(workbook.create_sheet(SHEETS[1]), state, ranked, evidence)
    _write_methode(workbook.create_sheet(SHEETS[2]), state, ranked, header, stamp)
    _write_spalten(workbook.create_sheet(SHEETS[3]), state.definitions)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- Ergebnisse ------------------------------------------------------------------


def _ranked_fields(definitions: Sequence[FieldDefinition]) -> list[_Ranked]:
    by_key = {d.key: d for d in definitions}
    seeded = sorted(
        (d for d in definitions if d.priority_rank is not None and d.excel_column),
        key=lambda d: (d.priority_rank or 0, d.key),
    )
    ranked = [_Ranked(d.key, d.excel_column or d.key, d.label_de, d.description_de) for d in seeded]
    for key in RANKED_FIELD_KEYS:
        if key not in by_key or by_key[key].priority_rank is None:
            column, label = _FALLBACK_COLUMNS[key]
            ranked.append(_Ranked(key, column, label, None))
    return ranked


def _ergebnisse_header(ranked: Sequence[_Ranked]) -> list[str]:
    header = list(_LEADING)
    for f in ranked:
        header.append(f.column)
        if f.key == "municipality":
            header.append(f"{f.column}_BFS")
        header += [f"{f.column}_Konfidenz", f"{f.column}_Quellen"]
    return header + list(_TRAILING)


def _row(state: AssetState, ranked: Sequence[_Ranked], stamp: str) -> dict[str, Any]:
    asset = state.asset
    row: dict[str, Any] = {
        "Dateiname": asset.external_ref,
        "Datei_SHA256": state.master_sha256,
        "Prio": asset.priority,
        "QC_Status": state.qc,
    }
    if state.head is None:
        # No finished run yet: the row exists so the count matches the
        # delivery, and says nothing else (data-model §8.12).
        row["QC_Status"] = QC_OPEN
        return row

    row["Signatur"] = asset.signature
    emitted: list[tuple[_Ranked, RecordField]] = []
    review: list[str] = []
    undetermined: list[str] = []
    for f in ranked:
        rf = state.fields.get(f.key)
        if rf is None:
            continue
        if rf.decision == FieldDecisionKind.ABSTAIN:
            reason = str(rf.abstain_reason) if rf.abstain_reason else "insufficient_evidence"
            undetermined.append(f"{f.column} ({reason})")
            continue
        emitted.append((f, rf))
        if rf.decision == FieldDecisionKind.EMIT_WITH_REVIEW:
            review.append(f.column)
        row[f.column] = rf.value_text
        row[f"{f.column}_Konfidenz"] = _number(rf.confidence)
        row[f"{f.column}_Quellen"] = rf.explanation or _render_sources(rf)
        if f.key == "municipality" and isinstance(rf.value, dict):
            row[f"{f.column}_BFS"] = _int(rf.value.get("bfs_nr"))

    subject = _point(state.fields.get("subject_point"))
    if subject is not None:
        row["Motiv_LV95_E"], row["Motiv_LV95_N"] = subject
    # P0 resolves no camera pose. If a later stage stores one, it is printed;
    # until then the three columns stay empty rather than guessed.
    camera_field = state.fields.get("camera_point")
    camera = _point(camera_field)
    if camera is not None and camera_field is not None:
        row["Aufnahmestandort_LV95_E"], row["Aufnahmestandort_LV95_N"] = camera
        if isinstance(camera_field.value, dict):
            row["Aufnahmerichtung_Grad"] = _number(camera_field.value.get("azimuth_deg"), 1)

    if asset.captured_at is not None:
        row["Aufnahmedatum_EXIF"] = asset.captured_at.date()
    row["Ausserhalb_SZ"] = "ja" if asset.outside_canton else "nein"

    confidences = [c for _, rf in emitted if (c := _number(rf.confidence)) is not None]
    if confidences:
        # From the head version, so a human correction moves it; the run's
        # `outcome_confidence` is the same number for an untouched record.
        row["Gesamtkonfidenz"] = min(confidences)
    elif emitted and state.run is not None:
        row["Gesamtkonfidenz"] = _number(state.run.outcome_confidence)
    if not emitted:
        row["Status"] = STATUS_NONE
    elif all(rf.decision == FieldDecisionKind.EMIT for _, rf in emitted):
        row["Status"] = STATUS_ACCEPT
    else:
        row["Status"] = STATUS_REVIEW
    row["Pruefbedarf"] = ", ".join(review) or None
    row["Unbestimmt"] = ", ".join(undetermined) or None
    row["QC_Von"] = state.last_reviewer if state.reviews else None
    if state.record is not None and state.record.publication_hold:
        row["Bemerkung"] = "Publikation zurückgehalten (Schutzvermerk)"
    row["Run_ID"] = state.run.public_id if state.run is not None else None
    row["Erzeugt_am"] = stamp
    return row


_WIDTHS: Mapping[str, float] = {
    "Dateiname": 22,
    "Datei_SHA256": 18,
    "Signatur": 14,
    "Prio": 6,
    "Gesamtkonfidenz": 10,
    "Status": 15,
    "Pruefbedarf": 22,
    "Unbestimmt": 34,
    "QC_Status": 12,
    "QC_Von": 18,
    "Bemerkung": 24,
    "Run_ID": 16,
    "Erzeugt_am": 21,
    "Aufnahmedatum_EXIF": 13,
    "Ausserhalb_SZ": 8,
}


def _width(column: str) -> float:
    if column in _WIDTHS:
        return _WIDTHS[column]
    if column.endswith("_Quellen"):
        return 48
    if column.endswith(("_Konfidenz", "_BFS", "_Grad")):
        return 10
    if "_LV95_" in column:
        return 11
    return 24


def _write_ergebnisse(ws: Worksheet, header: Sequence[str], rows: Sequence[dict[str, Any]]) -> None:
    _write_header(ws, header)
    formats = {c: "0.000" for c in header if c.endswith("_Konfidenz") or c == "Gesamtkonfidenz"} | {
        c: "0" for c in header if "_LV95_" in c or c.endswith("_BFS")
    }
    formats |= {"Aufnahmerichtung_Grad": "0.0", "Aufnahmedatum_EXIF": "yyyy-mm-dd"}
    status_col = header.index("Status") + 1
    for r, values in enumerate(rows, start=2):
        for c, column in enumerate(header, start=1):
            cell = _put(ws, r, c, values.get(column))
            if column in formats and cell.value is not None:
                cell.number_format = formats[column]
        status = values.get("Status")
        if isinstance(status, str) and status in _STATUS_STYLE:
            fill, font = _STATUS_STYLE[status]
            cell = ws.cell(row=r, column=status_col)
            cell.fill = fill
            cell.font = font
    for c, column in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(c)].width = _width(column)
    # Freeze the header and the file name: 36 columns scroll sideways, and the
    # row must stay identifiable while they do.
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(1, len(rows) + 1)}"


# --- Belege ----------------------------------------------------------------------


async def _cited_evidence(
    session: AsyncSession, state: BatchState
) -> dict[uuid.UUID, list[EvidenceItem]]:
    """Evidence cited by the emitted fields of each head version, by field id."""
    emitted = {
        rf.id
        for a in state.assets
        for rf in a.fields.values()
        if rf.decision != FieldDecisionKind.ABSTAIN
    }
    cited: dict[uuid.UUID, list[EvidenceItem]] = {}
    ids = sorted(emitted)
    for start in range(0, len(ids), 500):
        chunk = ids[start : start + 500]
        rows = await session.execute(
            select(RecordFieldEvidence.record_field_id, EvidenceItem)
            .join(EvidenceItem, EvidenceItem.id == RecordFieldEvidence.evidence_id)
            .where(RecordFieldEvidence.record_field_id.in_(chunk))
        )
        for field_id, item in rows:
            cited.setdefault(field_id, []).append(item)
    for items in cited.values():
        items.sort(key=lambda e: (-float(e.strength), e.local_id))
    return cited


def _write_belege(
    ws: Worksheet,
    state: BatchState,
    ranked: Sequence[_Ranked],
    evidence: Mapping[uuid.UUID, list[EvidenceItem]],
) -> None:
    _write_header(ws, _BELEGE_HEADER)
    columns = {f.key: f.column for f in ranked} | {
        d.key: d.excel_column or d.key for d in state.definitions if d.priority_rank is None
    }
    order = [f.key for f in ranked]
    r = 2
    for a in state.assets:
        keys = sorted(a.fields, key=lambda k: (order.index(k) if k in order else len(order), k))
        for key in keys:
            for item in evidence.get(a.fields[key].id, []):
                values = (
                    a.asset.external_ref,
                    columns.get(key, key),
                    _source_class_de(item.source_class),
                    item.source,
                    item.source_uri,
                    item.observation,
                    _number(item.strength),
                    str(item.agent_id) if item.agent_id else None,
                    item.model_id,
                )
                for c, value in enumerate(values, start=1):
                    cell = _put(ws, r, c, value)
                    if c == 7 and value is not None:
                        cell.number_format = "0.000"
                    if c == 6:
                        cell.alignment = _WRAP
                url = item.source_uri or ""
                if url.lower().startswith(("https://", "http://")):
                    link = ws.cell(row=r, column=5)
                    link.hyperlink = url
                    link.font = _LINK_FONT
                r += 1
    for c, width in enumerate((22, 16, 16, 30, 40, 70, 8, 14, 26), start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_BELEGE_HEADER))}{max(1, r - 1)}"


# --- Methode ---------------------------------------------------------------------


def _write_methode(
    ws: Worksheet,
    state: BatchState,
    ranked: Sequence[_Ranked],
    header: Sequence[str],
    stamp: str,
) -> None:
    accept, review = f"{ACCEPT_AT:.2f}", f"{REVIEW_AT:.2f}"
    single_model = WEIGHTS[SourceClass.VISUAL] * MODEL_STRENGTH_CAP
    lines = _Lines(ws)

    lines.title("Methode: Wie die Ortsangaben belegt und bewertet werden")
    lines.text(
        f"Export erstellt am {stamp} · Stapel «{state.batch.label or state.batch.id}» · "
        f"{len(state.assets)} Fotos, eine Zeile pro Foto im Blatt «Ergebnisse»."
    )
    lines.text(
        "Jede Angabe stützt sich auf Belege, die im Blatt «Belege» einzeln aufgeführt sind: "
        "Quelle, Link, Beobachtung und Stärke. Die Konfidenz ist keine Selbsteinschätzung eines "
        "Sprachmodells. Sie wird nach einer festen Regel aus diesen Belegen berechnet: Dieselben "
        "Belege ergeben immer dieselbe Zahl, und jede Zahl lässt sich von Hand nachrechnen."
    )

    lines.section("1. Quellenarten zählen, nicht Treffer")
    lines.text(
        "Die Aufgabenstellung nennt als Beispiel «zehn bestätigende Quellen = 100 %». Die Regel "
        "übernimmt diesen Gedanken: Je mehr voneinander unabhängige Bestätigungen, desto "
        "sicherer die Angabe. Sie macht ihn in zwei Punkten strenger."
    )
    lines.text(
        "Erstens zählt nicht die Zahl der Treffer, sondern die Zahl der unabhängigen "
        "Quellenarten. Zehn Webseiten, die voneinander abschreiben, sind eine Quelle und nicht "
        "zehn. Innerhalb einer Quellenart zählt nur der stärkste Beleg."
    )
    lines.text(
        "Zweitens kostet ein Widerspruch etwas. Ein anerkannter Einwand senkt die Konfidenz, "
        "eine Widerlegung setzt sie auf null."
    )

    lines.section("2. Die Formel")
    lines.text(
        f"raw = 1 {_MINUS} Π (1 {_MINUS} Gewicht {_TIMES} Stärke), "
        "über alle Quellenarten, die den Wert bestätigen"
    )
    lines.text(
        "Jede bestätigende Quellenart lässt einen Rest an Zweifel übrig, nämlich "
        f"(1 {_MINUS} Gewicht {_TIMES} Stärke). Diese Reste werden miteinander multipliziert, "
        "und was übrig "
        "bleibt, wird von 1 abgezogen. Eine weitere unabhängige Bestätigung erhöht die Konfidenz "
        "also immer, 100 % werden aber nie ganz erreicht."
    )
    for label, support in _examples():
        result = score(support)
        terms = f" {_TIMES} ".join(
            f"(1 {_MINUS} {WEIGHTS[s.source_class]:.2f} {_TIMES} {s.strength:.2f})" for s in support
        )
        lines.pair(
            label,
            f"1 {_MINUS} {terms} = {result.raw:.3f} → {_decision_de(result.decision)}",
        )

    lines.section("3. Quellenarten und ihr Gewicht")
    lines.row(("Quellenart", "Was darunter fällt", "Gewicht"), bold=True)
    for source_class, weight in sorted(WEIGHTS.items(), key=lambda kv: (-kv[1], kv[0].value)):
        lines.row((SOURCE_CLASS_DE[source_class], _SOURCE_CLASS_SCOPE[source_class], weight))

    lines.section("4. Wie die Stärke eines Belegs festgelegt wird")
    lines.text(
        "Die Stärke ist nicht frei wählbar. Werkzeuge, die amtliche Verzeichnisse abfragen, "
        "erhalten sie nach der Art der Übereinstimmung; Aussagen eines Modells sind gedeckelt."
    )
    lines.row(("Art des Belegs", "Beispiel", "Stärke"), bold=True)
    for label, example, strength in (
        (
            "Exakte Übereinstimmung",
            "Das amtliche Verzeichnis führt genau diesen Namen.",
            strength_for_match(MatchKind.EXACT),
        ),
        (
            "Übereinstimmung nach Normalisierung",
            "Gleich nach Angleichen von Gross-/Kleinschreibung, Akzenten, Bindestrichen "
            "und «St./Sankt».",
            strength_for_match(MatchKind.NORMALISED),
        ),
        (
            "Unscharfe Übereinstimmung",
            "Ähnlich, aber nicht gleich (Rangfolge einer Suchmaschine, Teiltreffer im Titel).",
            strength_for_match(MatchKind.FUZZY),
        ),
        (
            "Räumliche Abfrage",
            "Welche Ortschaft oder Gemeinde liegt an diesem Punkt?",
            strength_for_match(MatchKind.SPATIAL),
        ),
        (
            "Referenzbild mit passendem Titel",
            "Ein Vergleichsbild, dessen Titel dasselbe Objekt nennt.",
            REFERENCE_TITLE_MATCH,
        ),
        ("Webseite", "Je Seite; mehrere Seiten bleiben eine Quellenart.", OPEN_WEB_PAGE),
        (
            "Aussage eines Modells (höchstens)",
            "Was ein Sprachmodell im Bild sieht oder liest.",
            MODEL_STRENGTH_CAP,
        ),
    ):
        lines.row((label, example, strength))
    lines.text(
        f"Eine einzelne Modellbeobachtung allein kommt damit auf höchstens "
        f"{WEIGHTS[SourceClass.VISUAL]:.2f} {_TIMES} {MODEL_STRENGTH_CAP:.2f} "
        f"= {single_model:.2f}. "
        "Sie kann eine Angabe nie allein übernahmefähig machen."
    )

    lines.section("5. Einwände und Widerlegungen")
    lines.text(
        "Ein eigener Prüfschritt sucht gezielt nach Gründen, warum eine Angabe falsch sein "
        "könnte. Jeder anerkannte Einwand mit der Schwere s multipliziert die Konfidenz mit "
        f"(1 {_MINUS} s). Eine Widerlegung (Schwere 1 oder ein widersprechender Beleg) "
        "setzt die "
        "Konfidenz auf 0: Das Feld bleibt leer und erscheint in «Unbestimmt» mit dem Grund "
        "conflicting_evidence."
    )

    lines.section("6. Schwellen")
    lines.text(
        f"Konfidenz ≥ {accept} = übernahmefähig; darunter Prüfbedarf; unter {review} bleibt das "
        "Feld leer: unbekannte Angaben werden nicht geraten."
    )
    lines.pair("Übernahmefähig", f"Konfidenz ≥ {accept}")
    lines.pair("Prüfbedarf", f"{review} ≤ Konfidenz < {accept}; Wert steht da, bitte prüfen")
    lines.pair("Bewusst leer", f"Konfidenz < {review} oder widerlegt; kein Wert, keine Zahl")
    for code, meaning in _ABSTAIN_DE.items():
        lines.pair(f"Grund {code}", meaning)

    lines.section("7. Gesamtkonfidenz und Status einer Zeile")
    lines.text(
        "Gesamtkonfidenz ist die Konfidenz des schwächsten ausgegebenen Feldes: Eine Zeile ist "
        "nur so sicher wie ihre unsicherste Angabe. Leer gelassene Felder zählen dabei nicht "
        "mit; sie stehen in «Unbestimmt»."
    )
    lines.text(
        f"Status ist das Urteil der Regel: «{STATUS_ACCEPT}», wenn jedes ausgegebene Feld "
        f"mindestens {accept} erreicht; «{STATUS_REVIEW}», wenn mindestens eines darunter liegt; "
        f"«{STATUS_NONE}», wenn kein Feld ausgegeben wurde. QC_Status zeigt dagegen, was ein "
        "Mensch daraus gemacht hat: maschinell (noch niemand hat die Zeile angesehen), geprüft "
        "(jedes ausgegebene Feld unverändert bestätigt), korrigiert (mindestens ein Feld "
        "geändert, verworfen oder ergänzt), offen (Prüfung begonnen, aber nicht abgeschlossen, "
        "oder die Analyse läuft noch). Eine Korrektur geht vor einer Bestätigung."
    )

    lines.section("8. Kalibrierung")
    lines.text(
        "nicht kalibriert: die Zahl ist eine Regel über die Belege, keine gemessene Trefferquote."
    )
    lines.text(
        "Die Aufgabenstellung weist selbst darauf hin, dass eine selbst eingeschätzte Konfidenz "
        "keine gemessene Trefferquote ist. Eine Kalibrierung an geprüften Fotos ist vorgesehen; "
        "bis dahin gilt Konfidenz = Rohwert der Regel."
    )

    lines.section("9. Die Spalten im Blatt «Ergebnisse»")
    lines.row(("Spalte", "Bedeutung", ""), bold=True)
    definitions = _column_definitions(ranked)
    for column in header:
        lines.pair(column, definitions.get(column, ""))

    lines.section("10. Verwendete Werte")
    lines.pair("Schwelle übernahmefähig", accept)
    lines.pair("Schwelle Prüfbedarf", review)
    lines.pair("Höchste Stärke einer Modellaussage", f"{MODEL_STRENGTH_CAP:.2f}")
    lines.pair("Kalibrierung", "keine (Konfidenz = Rohwert)")
    lines.pair("Export erstellt am", stamp)

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 100
    ws.column_dimensions["C"].width = 10


def _examples() -> list[tuple[str, list[Support]]]:
    exact = strength_for_match(MatchKind.EXACT)
    geodata = Support(SourceClass.GEODATA_LAYER, exact)
    reference = Support(SourceClass.REFERENCE_IMAGE, REFERENCE_TITLE_MATCH)
    text = Support(SourceClass.TEXT_IN_IMAGE, MODEL_STRENGTH_CAP)
    visual = Support(SourceClass.VISUAL, MODEL_STRENGTH_CAP)
    return [
        ("Beispiel: nur ein Modell", [visual]),
        ("Beispiel: Geodaten + Referenzbild", [geodata, reference]),
        ("Beispiel: + Schrift im Bild", [geodata, reference, text]),
    ]


def _decision_de(decision: FieldDecisionKind) -> str:
    return {
        FieldDecisionKind.EMIT: STATUS_ACCEPT,
        FieldDecisionKind.EMIT_WITH_REVIEW: STATUS_REVIEW,
        FieldDecisionKind.ABSTAIN: "bleibt leer",
    }[decision]


def _column_definitions(ranked: Sequence[_Ranked]) -> dict[str, str]:
    definitions = {
        "Dateiname": "Name der gelieferten Bilddatei; der eindeutige Schlüssel zur Datei.",
        "Datei_SHA256": "Prüfsumme der Bilddatei (Vorderseite, Master). Die Zuordnung bleibt "
        "auch nach einer Umbenennung eindeutig.",
        "Signatur": "Signatur des Staatsarchivs; leer, solange keine vergeben ist.",
        "Prio": "Priorität der Lieferung (1 = eines der Prio-1-Fotos).",
        "Motiv_LV95_E": "Ostwert (LV95) des Punktes, den das Foto zeigt; leer, wenn nicht "
        "bestimmt.",
        "Motiv_LV95_N": "Nordwert (LV95) des Punktes, den das Foto zeigt; leer, wenn nicht "
        "bestimmt.",
        "Aufnahmestandort_LV95_E": "Ostwert (LV95) des Kamerastandorts. Wird noch nicht "
        "bestimmt und bleibt deshalb leer, statt geschätzt zu werden.",
        "Aufnahmestandort_LV95_N": "Nordwert (LV95) des Kamerastandorts; wie oben leer.",
        "Aufnahmerichtung_Grad": "Blickrichtung der Kamera in Grad ab Nord; wie oben leer.",
        "Aufnahmedatum_EXIF": "Aufnahmedatum aus den EXIF-Daten der Datei; leer, wenn die Datei "
        "keines enthält.",
        "Ausserhalb_SZ": "ja, wenn das Motiv ausserhalb des Kantons Schwyz liegt.",
        "Gesamtkonfidenz": "Konfidenz des schwächsten ausgegebenen Feldes (Abschnitt 7).",
        "Status": f"{STATUS_ACCEPT} · {STATUS_REVIEW} · {STATUS_NONE}: das Urteil der Regel "
        "über die ganze Zeile (Abschnitt 7).",
        "Pruefbedarf": "Felder, deren Wert unter der Übernahmeschwelle liegt und geprüft "
        "werden sollte.",
        "Unbestimmt": "Bewusst leer gelassene Felder, mit Grund (Abschnitt 6).",
        "QC_Status": "maschinell · geprüft · korrigiert · offen: der Stand der menschlichen "
        "Prüfung (Abschnitt 7).",
        "QC_Von": "Wer die Zeile zuletzt geprüft oder korrigiert hat.",
        "Bemerkung": "Hinweis des Archivs, etwa ein Schutzvermerk.",
        "Run_ID": "Kennung der Analyse, die die Werte erzeugt hat; im Programm nachschlagbar.",
        "Erzeugt_am": "Zeitpunkt dieses Exports (UTC).",
    }
    for f in ranked:
        about = f" {f.description}" if f.description else ""
        definitions[f.column] = f"{f.label}.{about} Leer, wenn bewusst leer gelassen."
        definitions[f"{f.column}_Konfidenz"] = (
            "Konfidenz nach der Regel (Abschnitte 2 bis 6), drei Dezimalen; leer, wenn das Feld "
            "bewusst leer gelassen wurde, nie 0."
        )
        definitions[f"{f.column}_Quellen"] = (
            "Welche Quellenarten den Wert bestätigen und wie viele Einwände die Prüfung "
            "anerkannt hat; die einzelnen Belege stehen im Blatt «Belege»."
        )
        if f.key == "municipality":
            definitions[f"{f.column}_BFS"] = "BFS-Gemeindenummer der politischen Gemeinde."
    return definitions


class _Lines:
    """A cursor over the Methode sheet: title, sections, paragraphs, tables."""

    def __init__(self, ws: Worksheet) -> None:
        self.ws = ws
        self.r = 1

    def title(self, text: str) -> None:
        _put(self.ws, self.r, 1, text).font = _TITLE_FONT
        self.r += 2

    def section(self, text: str) -> None:
        self.r += 1
        _put(self.ws, self.r, 1, text).font = _SECTION_FONT
        self.r += 1

    def text(self, text: str) -> None:
        # A paragraph spans A:B so it reads as prose, not as a label.
        self.ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=2)
        _put(self.ws, self.r, 1, text).alignment = _WRAP
        self.ws.row_dimensions[self.r].height = max(15.0, 15.0 * (len(text) // 120 + 1))
        self.r += 1

    def pair(self, label: str, text: str) -> None:
        _put(self.ws, self.r, 1, label).alignment = _TOP
        _put(self.ws, self.r, 2, text).alignment = _WRAP
        self.r += 1

    def row(self, values: Sequence[Any], *, bold: bool = False) -> None:
        for c, value in enumerate(values, start=1):
            cell = _put(self.ws, self.r, c, value if value != "" else None)
            cell.alignment = _WRAP
            if bold:
                cell.font = _HEADER_FONT
            if isinstance(value, float):
                cell.number_format = "0.00"
        self.r += 1


# --- Spalten ---------------------------------------------------------------------


def _write_spalten(ws: Worksheet, definitions: Sequence[FieldDefinition]) -> None:
    header = (
        "Excel-Spalte",
        "Feld (field_key)",
        "Bezeichnung",
        "Rang",
        "ISAD(G)",
        "Dublin Core",
        "EAD3",
        "scopeArchiv",
        "Beschreibung",
    )
    _write_header(ws, header)
    ordered = sorted(
        definitions, key=lambda d: (d.priority_rank is None, d.priority_rank or 0, d.key)
    )
    for r, d in enumerate(ordered, start=2):
        for c, value in enumerate(
            (
                d.excel_column,
                d.key,
                d.label_de,
                d.priority_rank,
                d.isad,
                d.dublin_core,
                d.ead3,
                d.scope_archiv,
                d.description_de,
            ),
            start=1,
        ):
            _put(ws, r, c, value).alignment = _WRAP if c == 9 else _TOP
    note_row = len(ordered) + 3
    _put(
        ws,
        note_row,
        1,
        "Die Zuordnung ist Daten, nicht Programmcode: Ein späteres Importformat (etwa "
        "scopeArchiv) ist eine Änderung in dieser Tabelle. Leere Zellen sind noch nicht "
        "festgelegt.",
    )
    for c, width in enumerate((24, 20, 24, 6, 10, 18, 34, 18, 70), start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"


# --- cells -------------------------------------------------------------------------


def _write_header(ws: Worksheet, header: Sequence[str]) -> None:
    for c, column in enumerate(header, start=1):
        cell = ws.cell(row=1, column=c, value=column)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")


def _put(ws: Worksheet, row: int, column: int, value: Any) -> Any:
    """Write a value. Strings are stored as strings, never as formulas."""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if value == "":
            value = None
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def _number(value: Decimal | float | int | str | None, places: int = 3) -> float | None:
    if value is None:
        return None
    try:
        return round(float(value), places)
    except (TypeError, ValueError):
        return None


def _int(value: Any) -> int | None:
    try:
        return int(value) if value is not None else None
    except (TypeError, ValueError):
        return None


def _point(rf: RecordField | None) -> tuple[float, float] | None:
    """`{lv95: [e, n]}` of an emitted point field, else None."""
    if rf is None or rf.decision == FieldDecisionKind.ABSTAIN or not isinstance(rf.value, dict):
        return None
    lv95 = rf.value.get("lv95")
    if not isinstance(lv95, list | tuple) or len(lv95) != 2:
        return None
    e, n = _number(lv95[0], 1), _number(lv95[1], 1)
    return (e, n) if e is not None and n is not None else None


def _source_class_de(value: SourceClass | str) -> str:
    try:
        return SOURCE_CLASS_DE[SourceClass(value)]
    except ValueError:
        return str(value)


def _render_sources(rf: RecordField) -> str:
    """The source line when a version predates the stored `explanation`."""
    names = [_source_class_de(c) for c in rf.source_classes]
    n, k = len(names), rf.objection_count
    einwaende = "1 anerkannter Einwand" if k == 1 else f"{k} anerkannte Einwände"
    if n == 0:
        return f"{rf.corroborating_count} Belege · {einwaende}"
    verb = "Quellenart bestätigt" if n == 1 else "Quellenarten bestätigen"
    return f"{n} {verb} ({', '.join(names)}) · {einwaende}"


def _iso(moment: datetime) -> str:
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=UTC)
    return moment.astimezone(UTC).isoformat(timespec="seconds").replace("+00:00", "Z")
