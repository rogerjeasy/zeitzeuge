"""Withdraw, restore, bulk actions and the export selection (batch-management note).

Built on the end-to-end fixtures of `test_investigation_api`: two photographs
are uploaded against moto, one is given a machine record, and the actions are
exercised through the HTTP API exactly as the web app calls them.
"""

# ruff: noqa: F811  (pytest fixtures imported from test_investigation_api)
from __future__ import annotations

import io

import boto3
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from tests.conftest import BUCKET, REGION
from tests.test_investigation_api import _machine_run, investigation  # noqa: F401  (fixture)

from zeitzeuge_api.db.models import Run, User
from zeitzeuge_api.db.models.enums import RunStatus, UserRole
from zeitzeuge_api.db.session import TenantContext


async def _batch_with_two(client) -> tuple[str, list[dict]]:
    batch_id = (await client.post("/api/v1/batches", json={"label": "prio1"})).json()["id"]
    files = (
        await client.post(
            f"/api/v1/batches/{batch_id}/assets",
            json={
                "files": [
                    {"filename": name, "content_type": "image/jpeg", "content_length": 10}
                    for name in ("SG.CIV.17.3.jpg", "SG.CIV.17.11.jpg")
                ]
            },
        )
    ).json()["files"]
    s3 = boto3.client("s3", region_name=REGION)
    committed = []
    for file in files:
        s3.put_object(
            Bucket=BUCKET,
            Key=file["upload"]["fields"]["key"],
            Body=b"x" * 10,
            ContentType="image/jpeg",
        )
        committed.append(
            {**file, **(await client.put(f"/api/v1/assets/{file['asset_id']}/commit")).json()}
        )
    return batch_id, committed


async def _set_role(app, email: str, role: UserRole) -> None:
    async with app.state.database.session(TenantContext.bootstrap()) as session:
        user = await session.scalar(select(User).where(User.email == email))
        assert user is not None
        user.role = role


async def test_withdraw_hides_cancels_and_restore_requeues(
    investigation,
    client,
    as_user,
) -> None:
    await _set_role(investigation, as_user["email"], UserRole.ARCHIVIST)
    batch_id, (first, _second) = await _batch_with_two(client)

    withdrawn = await client.request(
        "DELETE", f"/api/v1/assets/{first['asset_id']}", json={"reason": "Doppel"}
    )
    assert withdrawn.status_code == 200, withdrawn.text
    assert withdrawn.json()["succeeded"] == 1

    batch = (await client.get(f"/api/v1/batches/{batch_id}")).json()
    assert [a["external_ref"] for a in batch["assets"]] == ["SG.CIV.17.11.jpg"]
    assert batch["counts"]["total"] == 1 and batch["counts"]["deleted"] == 1
    trash = (await client.get(f"/api/v1/batches/{batch_id}?include_deleted=true")).json()
    gone = next(a for a in trash["assets"] if a["asset_id"] == first["asset_id"])
    assert gone["deleted_at"] and gone["deletion_reason"] == "Doppel"
    assert gone["deleted_by_name"]  # the Papierkorb says who
    assert trash["counts"]["total"] == 1  # listed, never counted

    async with investigation.state.database.session(TenantContext(role="worker")) as session:
        run = await session.scalar(select(Run).where(Run.public_id == first["run_id"]))
        assert run is not None and run.status is RunStatus.CANCELLED

    again = await client.request("DELETE", f"/api/v1/assets/{first['asset_id']}")
    assert again.json()["results"][0]["detail"] == "already_withdrawn"

    restored = await client.post(f"/api/v1/assets/{first['asset_id']}/restore")
    assert restored.json()["succeeded"] == 1
    batch = (await client.get(f"/api/v1/batches/{batch_id}")).json()
    assert batch["counts"]["total"] == 2 and batch["counts"]["deleted"] == 0
    assert batch["counts"]["queued"] == 2  # the cancelled run was replaced by a fresh one


async def test_bulk_accept_confirms_every_open_field_in_one_version(
    investigation,
    client,
    as_user,
) -> None:
    _batch_id, (first, second) = await _batch_with_two(client)
    await _machine_run(investigation, first["run_id"])
    before = (await client.get(f"/api/v1/assets/{first['asset_id']}/case")).json()

    result = await client.post(
        "/api/v1/assets/bulk",
        json={"action": "accept", "asset_ids": [first["asset_id"], second["asset_id"]]},
    )
    assert result.status_code == 200, result.text
    by_id = {r["asset_id"]: r for r in result.json()["results"]}
    assert by_id[first["asset_id"]]["ok"] is True
    assert by_id[second["asset_id"]]["detail"] == "no_record"

    after = (await client.get(f"/api/v1/assets/{first['asset_id']}/case")).json()
    assert after["version"] == before["version"] + 1
    assert after["qc_status"] == "geprüft"
    emitted = [f for f in after["fields"] if f["rank"] and f["decision"] != "abstain"]
    assert emitted and all(f["review_action"] == "accepted_unchanged" for f in emitted)
    empty = [f for f in after["fields"] if f["rank"] and f["decision"] == "abstain"]
    assert all(f["review_action"] is None for f in empty)  # nothing invented for empty fields

    twice = await client.post(
        "/api/v1/assets/bulk", json={"action": "accept", "asset_ids": [first["asset_id"]]}
    )
    assert twice.json()["results"][0]["detail"] == "nothing_to_accept"


async def test_export_selection_and_withdrawn_photos_are_left_out(
    investigation,
    client,
    as_user,
) -> None:
    await _set_role(investigation, as_user["email"], UserRole.ARCHIVIST)
    batch_id, (_first, second) = await _batch_with_two(client)

    def rows(content: bytes) -> list[str]:
        sheet = load_workbook(io.BytesIO(content))["Ergebnisse"]
        return [r[0].value for r in sheet.iter_rows(min_row=2)]

    selected = await client.get(
        f"/api/v1/batches/{batch_id}/export?fmt=xlsx&asset_ids={second['asset_id']}"
    )
    assert rows(selected.content) == ["SG.CIV.17.11.jpg"]

    await client.request("DELETE", f"/api/v1/assets/{second['asset_id']}")
    whole = await client.get(f"/api/v1/batches/{batch_id}/export?fmt=xlsx")
    assert rows(whole.content) == ["SG.CIV.17.3.jpg"]


async def test_a_batch_withdraws_and_restores_with_its_photos(
    investigation,
    client,
    as_user,
) -> None:
    await _set_role(investigation, as_user["email"], UserRole.ARCHIVIST)
    batch_id, (first, _second) = await _batch_with_two(client)
    # One photo withdrawn on its own first: restoring the batch must not bring it back.
    await client.request("DELETE", f"/api/v1/assets/{first['asset_id']}")

    gone = await client.request("DELETE", f"/api/v1/batches/{batch_id}")
    assert gone.status_code == 200 and gone.json()["assets"] == 1
    assert batch_id not in [b["id"] for b in (await client.get("/api/v1/batches")).json()]
    listed = (await client.get("/api/v1/batches?include_deleted=true")).json()
    assert next(b for b in listed if b["id"] == batch_id)["deleted_at"]

    back = await client.post(f"/api/v1/batches/{batch_id}/restore")
    assert back.json()["assets"] == 1
    batch = (await client.get(f"/api/v1/batches/{batch_id}")).json()
    assert [a["external_ref"] for a in batch["assets"]] == ["SG.CIV.17.11.jpg"]


@pytest.mark.parametrize("role", [UserRole.REVIEWER, UserRole.OPERATOR])
async def test_only_archivists_withdraw_but_reviewers_may_rerun(
    investigation,
    client,
    as_user,
    role: UserRole,
) -> None:
    batch_id, (first, _) = await _batch_with_two(client)
    await _set_role(investigation, as_user["email"], role)

    denied = await client.request("DELETE", f"/api/v1/assets/{first['asset_id']}")
    assert denied.status_code == 403
    denied_batch = await client.request("DELETE", f"/api/v1/batches/{batch_id}")
    assert denied_batch.status_code == 403
    rerun = await client.post(
        "/api/v1/assets/bulk", json={"action": "rerun", "asset_ids": [first["asset_id"]]}
    )
    assert rerun.status_code == 200
    assert rerun.json()["results"][0]["detail"] == "running"  # still queued from the upload


async def test_a_withdrawn_photo_cannot_be_committed_reviewed_or_accepted(
    investigation, client, as_user
) -> None:
    await _set_role(investigation, as_user["email"], UserRole.ARCHIVIST)
    _batch_id, (first, _second) = await _batch_with_two(client)
    await _machine_run(investigation, first["run_id"])
    case = (await client.get(f"/api/v1/assets/{first['asset_id']}/case")).json()
    await client.request("DELETE", f"/api/v1/assets/{first['asset_id']}")

    assert (await client.put(f"/api/v1/assets/{first['asset_id']}/commit")).status_code == 409
    review = await client.patch(
        f"/api/v1/records/{case['record_id']}/fields/locality", json={"action": "accept"}
    )
    assert review.status_code == 409
    bulk = await client.post(
        "/api/v1/assets/bulk", json={"action": "accept", "asset_ids": [first["asset_id"]]}
    )
    assert bulk.json()["results"][0]["detail"] == "withdrawn"


async def test_an_operator_does_not_decide_on_fields(investigation, client, as_user) -> None:
    _batch_id, (first, _second) = await _batch_with_two(client)
    await _machine_run(investigation, first["run_id"])
    case = (await client.get(f"/api/v1/assets/{first['asset_id']}/case")).json()
    await _set_role(investigation, as_user["email"], UserRole.OPERATOR)

    bulk = await client.post(
        "/api/v1/assets/bulk", json={"action": "accept", "asset_ids": [first["asset_id"]]}
    )
    assert bulk.status_code == 403
    review = await client.patch(
        f"/api/v1/records/{case['record_id']}/fields/locality", json={"action": "accept"}
    )
    assert review.status_code == 403


async def test_the_papierkorb_is_for_approvers_and_names_no_address(
    investigation, client, as_user
) -> None:
    await _set_role(investigation, as_user["email"], UserRole.ARCHIVIST)
    batch_id, (first, _second) = await _batch_with_two(client)
    await client.request("DELETE", f"/api/v1/assets/{first['asset_id']}", json={"reason": "x"})
    trash = (await client.get(f"/api/v1/batches/{batch_id}?include_deleted=true")).json()
    gone = next(a for a in trash["assets"] if a["asset_id"] == first["asset_id"])
    assert "@" not in (gone["deleted_by_name"] or "")

    await _set_role(investigation, as_user["email"], UserRole.REVIEWER)
    assert (await client.get(f"/api/v1/batches/{batch_id}?include_deleted=true")).status_code == 403
    assert (await client.get("/api/v1/batches?include_deleted=true")).status_code == 403
    case = (await client.get(f"/api/v1/assets/{first['asset_id']}/case")).json()
    assert case["deleted_at"] and case["deletion_reason"] is None


async def test_timing_reaches_the_batch_the_case_and_the_workbook(
    investigation, client, as_user
) -> None:
    from datetime import UTC, datetime, timedelta

    batch_id, (first, _second) = await _batch_with_two(client)
    async with investigation.state.database.session(TenantContext(role="worker")) as session:
        run = await session.scalar(select(Run).where(Run.public_id == first["run_id"]))
        assert run is not None
        run.started_at = datetime.now(UTC) - timedelta(seconds=95)
    await _machine_run(investigation, first["run_id"])

    batch = (await client.get(f"/api/v1/batches/{batch_id}")).json()
    row = next(a for a in batch["assets"] if a["asset_id"] == first["asset_id"])
    assert row["queued_at"] and row["started_at"] and row["finished_at"]
    assert 90_000 <= row["duration_ms"] <= 120_000
    assert batch["timing"]["finished"] == 1
    assert batch["timing"]["median_ms"] == row["duration_ms"]

    case = (await client.get(f"/api/v1/assets/{first['asset_id']}/case")).json()
    assert case["run"]["duration_ms"] == row["duration_ms"]
    assert [s["stage"] for s in case["run"]["stages"]] == ["grounding", "hypothesise"]
    assert case["run"]["stages"][0]["duration_ms"] is not None

    sheet = load_workbook(
        io.BytesIO((await client.get(f"/api/v1/batches/{batch_id}/export?fmt=xlsx")).content)
    )["Ergebnisse"]
    header = [c.value for c in sheet[1]]
    values = {
        r[0].value: dict(zip(header, [c.value for c in r], strict=False))
        for r in sheet.iter_rows(min_row=2)
    }
    took = values["SG.CIV.17.3.jpg"]["Untersuchung_Dauer_s"]
    assert took == round(row["duration_ms"] / 1000, 1)
    assert values["SG.CIV.17.3.jpg"]["Untersuchung_Beginn"].endswith("Z")
    assert values["SG.CIV.17.11.jpg"]["Untersuchung_Dauer_s"] is None  # not investigated yet


async def _fail(app, run_public_id: str) -> None:
    async with app.state.database.session(TenantContext(role="worker")) as session:
        run = await session.scalar(select(Run).where(Run.public_id == run_public_id))
        assert run is not None
        run.status = RunStatus.FAILED


async def test_failed_runs_are_deleted_for_good_and_finished_work_is_kept(
    investigation, client, as_user
) -> None:
    await _set_role(investigation, as_user["email"], UserRole.ARCHIVIST)
    batch_id, (first, second) = await _batch_with_two(client)
    await _fail(investigation, first["run_id"])
    # The second run wrote a record, then failed on a later attempt's bookkeeping:
    # a run a record version cites is never purged.
    await _machine_run(investigation, second["run_id"])
    await _fail(investigation, second["run_id"])

    purged = await client.delete(f"/api/v1/runs/failed?batch_id={batch_id}")
    assert purged.status_code == 200, purged.text
    assert purged.json()["purged"] == 1

    async with investigation.state.database.session(TenantContext(role="worker")) as session:
        left = {r.public_id for r in await session.scalars(select(Run))}
    assert first["run_id"] not in left
    assert second["run_id"] in left

    again = await client.delete("/api/v1/runs/failed")
    assert again.json()["purged"] == 0


async def test_only_archivists_purge_failed_runs(investigation, client, as_user) -> None:
    await _set_role(investigation, as_user["email"], UserRole.REVIEWER)
    refused = await client.delete("/api/v1/runs/failed")
    assert refused.status_code == 403


async def test_the_qc_endpoint_answers_with_the_measured_hit_rate(
    investigation, client, as_user
) -> None:
    batch_id, (first, _second) = await _batch_with_two(client)
    await _machine_run(investigation, first["run_id"])
    case = (await client.get(f"/api/v1/assets/{first['asset_id']}/case")).json()
    record_id, version = case["record_id"], case["version"]
    accepted = await client.patch(
        f"/api/v1/records/{record_id}/fields/locality",
        json={"action": "accept"},
        headers={"If-Match": str(version)},
    )
    assert accepted.status_code == 200, accepted.text

    qc = await client.get(f"/api/v1/batches/{batch_id}/qc")
    assert qc.status_code == 200, qc.text  # the schema forbids unknown keys
    measured = qc.json()["measured"]
    assert measured["judged"] == 1
    assert measured["fields"]["locality"]["correct"] == 1


async def test_history_collections_and_open_questions_show_real_work(
    investigation, client, as_user
) -> None:
    batch_id, (first, _second) = await _batch_with_two(client)
    await _machine_run(investigation, first["run_id"])
    case = (await client.get(f"/api/v1/assets/{first['asset_id']}/case")).json()
    await client.patch(
        f"/api/v1/records/{case['record_id']}/fields/locality",
        json={"action": "accept"},
        headers={"If-Match": str(case["version"])},
    )

    history = await client.get("/api/v1/history")
    assert history.status_code == 200, history.text
    actions = {(e["action"], e["photo"]) for e in history.json()}
    assert ("record.field_reviewed", "SG.CIV.17.3.jpg") in actions
    assert any(e["action"] == "asset.uploaded" for e in history.json())
    assert all("ip" not in key for e in history.json() for key in e)

    collections = (await client.get("/api/v1/collections")).json()
    fonds = next(c for c in collections if c["photos"])
    assert (fonds["photos"], fonds["investigated"], fonds["reviewed"]) == (2, 1, 1)
    assert fonds["described"] == 1

    questions = (await client.get("/api/v1/open-questions")).json()
    first_q = next(q for q in questions if q["photo"] == "SG.CIV.17.3.jpg")
    assert "building" in first_q["open_fields"]
    assert first_q["known"]["locality"] == "Schwyz"
    assert batch_id == first_q["batch_id"]


async def test_guests_do_not_see_the_history(investigation, client, as_user) -> None:
    await _set_role(investigation, as_user["email"], UserRole.GUEST)
    assert (await client.get("/api/v1/history")).status_code == 403


async def test_the_cost_report_splits_by_vendor_model_stage_and_batch(
    investigation, client, as_user
) -> None:
    import uuid as _uuid
    from datetime import UTC, datetime, timedelta
    from decimal import Decimal

    from zeitzeuge_api.db.models import AgentInvocation
    from zeitzeuge_api.db.models.enums import AgentId, InvocationStatus

    batch_id, (first, _second) = await _batch_with_two(client)
    await _machine_run(investigation, first["run_id"])
    async with investigation.state.database.session(TenantContext(role="worker")) as session:
        run = await session.scalar(select(Run).where(Run.public_id == first["run_id"]))
        assert run is not None
        start = datetime.now(UTC) - timedelta(minutes=5)

        def call(stage: str, model: str, cost: str, output: dict | None = None) -> AgentInvocation:
            return AgentInvocation(
                id=_uuid.uuid4(),
                org_id=run.org_id,
                run_id=run.id,
                agent_id=AgentId.SCENE,
                stage=stage,
                acting_as=["scene"],
                tier=1,
                status=InvocationStatus.SUCCEEDED,
                model_id=model,
                inference_profile=model,
                region="eu",
                output=output,
                started_at=start,
                ended_at=start + timedelta(seconds=30),
                input_tokens=1000,
                output_tokens=200,
                cost_usd=Decimal(cost),
            )

        session.add(call("perception", "eu.anthropic.claude-opus-4-6-v1", "0.0800"))
        session.add(
            call(
                "research",
                "gpt-5.5-2026-04-23, gemini-4-argon",
                "0.5000",
                {
                    "cost_by_model": [
                        {
                            "model_id": "gpt-5.5-2026-04-23",
                            "input_tokens": 900,
                            "output_tokens": 100,
                            "cost_usd": 0.45,
                        },
                        {
                            "model_id": "gemini-4-argon",
                            "input_tokens": 100,
                            "output_tokens": 100,
                            "cost_usd": 0.05,
                        },
                    ]
                },
            )
        )

    report = await client.get("/api/v1/costs?days=7")
    assert report.status_code == 200, report.text
    body = report.json()
    assert body["totals"]["cost_usd"] == 0.58
    labels = {m["label"]: m for m in body["by_model"]}
    assert labels["GPT-5.5 + Websuche"]["cost_usd"] == 0.45
    assert labels["Claude Opus 4.6"]["vendor"] == "anthropic"
    today = body["by_day"][-1]
    assert (today["anthropic"], today["openai"], today["google"]) == (0.08, 0.45, 0.05)
    assert {s["stage"] for s in body["by_stage"]} >= {"perception", "research"}
    assert body["by_batch"][0]["batch_id"] == batch_id
    assert "google_cloud_vision" in body["unmetered"]
    assert (await client.get("/api/v1/costs?days=3")).status_code in (400, 422)


async def test_guests_do_not_see_the_costs(investigation, client, as_user) -> None:
    await _set_role(investigation, as_user["email"], UserRole.GUEST)
    assert (await client.get("/api/v1/costs")).status_code == 403


def _labelled_workbook() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    green = PatternFill(start_color="FF92D050", end_color="FF92D050", fill_type="solid")
    book = Workbook()
    sheet = book.active
    sheet.append(["Dateiname", "Ortschaft", "Gemeinde", "Gebaeude", "See_Fluss_Berg", "Strasse"])
    sheet.append(["SG.CIV.17.3.jpg", "Schwyz", "Ingenbohl", "Kapelle Haggenegg", "Mythen", None])
    for column in ("B", "C", "D"):  # See_Fluss_Berg is not green: not a label
        sheet[f"{column}2"].fill = green
    sheet.append(["SG.CIV.17.999.jpg", "Arth", None, None, None, None])
    sheet["B3"].fill = green
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


async def test_the_teams_labelled_workbook_becomes_reviews_once(
    investigation, client, as_user
) -> None:
    batch_id, (first, _second) = await _batch_with_two(client)
    await _machine_run(investigation, first["run_id"])
    url = f"/api/v1/batches/{batch_id}/qc-import?filename=labels.xlsx"

    result = await client.post(url, content=_labelled_workbook())
    assert result.status_code == 200, result.text
    body = result.json()
    # Schwyz names the machine's Schwyz; Ingenbohl corrects it; the building was open.
    assert (body["accepted"], body["edited"], body["filled"]) == (1, 1, 1)
    assert body["unmatched"] == ["SG.CIV.17.999.jpg"]

    qc = (await client.get(f"/api/v1/batches/{batch_id}/qc")).json()
    rates = qc["overrideRate"]
    assert rates["locality"] == 0.0
    assert rates["municipality"] == 1.0
    assert rates["building"] == 1.0
    assert rates["landscape_feature"] is None  # an uncoloured cell is no label
    assert qc["measured"]["fields"]["municipality"]["wrong"] == 1

    again = (await client.post(url, content=_labelled_workbook())).json()
    assert again["already_reviewed"] == 3
    assert again["accepted"] + again["edited"] + again["filled"] == 0


async def test_a_file_that_is_not_a_results_sheet_is_refused(
    investigation, client, as_user
) -> None:
    batch_id, _ = await _batch_with_two(client)
    bad = await client.post(f"/api/v1/batches/{batch_id}/qc-import", content=b"not a workbook")
    assert bad.status_code == 422


async def test_a_zip_bomb_is_refused_before_parsing(investigation, client, as_user) -> None:
    import zipfile

    batch_id, _ = await _batch_with_two(client)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"\0" * (4 * 1024 * 1024))
    bomb = await client.post(f"/api/v1/batches/{batch_id}/qc-import", content=out.getvalue())
    assert bomb.status_code == 422
    assert "too large" in bomb.json()["detail"]
