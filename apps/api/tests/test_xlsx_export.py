"""The Excel workbook: the brief's mandatory deliverable (data-model §8.12).

One batch, three photos, each a state the archivist will meet in the real
delivery: a finished machine row with a field left open on purpose, a row a
human has confirmed, and a photo the agents have not reached yet. The workbook
is opened with openpyxl and read the way the archive will read it.
"""

from __future__ import annotations

import importlib.util
import io
import uuid
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from zeitzeuge_api.db.models import (
    Asset,
    AssetFile,
    CatalogueRecord,
    EvidenceItem,
    FieldDefinition,
    FieldReview,
    Fonds,
    Organization,
    RecordField,
    RecordFieldEvidence,
    RecordVersion,
    Run,
    UploadBatch,
    User,
)
from zeitzeuge_api.db.models.enums import (
    AbstainReason,
    ActorKind,
    AgentId,
    AssetSide,
    BatchSource,
    DescriptionLevel,
    EvidenceKind,
    ExtentUnit,
    FieldDecisionKind,
    FieldOrigin,
    FileRole,
    OrganizationKind,
    ReviewAction,
    RunStatus,
    SourceClass,
    VersionChange,
)
from zeitzeuge_api.db.session import TenantContext
from zeitzeuge_api.services.exporters.xlsx import build_workbook
from zeitzeuge_api.services.qc import qc_status, qc_summary

HEADER = [
    "Dateiname", "Datei_SHA256", "Signatur", "Prio",
    "Ortschaft", "Ortschaft_Konfidenz", "Ortschaft_Quellen",
    "Gemeinde", "Gemeinde_BFS", "Gemeinde_Konfidenz", "Gemeinde_Quellen",
    "Gebaeude", "Gebaeude_Konfidenz", "Gebaeude_Quellen",
    "See_Fluss_Berg", "See_Fluss_Berg_Konfidenz", "See_Fluss_Berg_Quellen",
    "Strasse", "Strasse_Konfidenz", "Strasse_Quellen",
    "Motiv_LV95_E", "Motiv_LV95_N",
    "Aufnahmestandort_LV95_E", "Aufnahmestandort_LV95_N", "Aufnahmerichtung_Grad",
    "Aufnahmedatum_EXIF", "Ausserhalb_SZ", "Gesamtkonfidenz", "Status",
    "Pruefbedarf", "Unbestimmt", "QC_Status", "QC_Von", "Bemerkung", "Run_ID",
    "Untersuchung_Beginn", "Untersuchung_Ende", "Untersuchung_Dauer_s", "Erzeugt_am",
]  # fmt: skip
"""§8.12 in order, with the one design-note addition `Status` after `Gesamtkonfidenz`."""

GENERATED_AT = datetime(2026, 10, 24, 9, 2, 44, tzinfo=UTC)
URL = "https://api3.geo.admin.ch/rest/services/api/MapServer/identify?layer=ortschaften"

MIGRATION = Path(__file__).resolve().parents[1] / "migrations" / "versions" / "0004_catalogue.py"


def _field_definitions() -> list[FieldDefinition]:
    """The seed of the shipping migration, loaded by path like `_policies.py`
    does, so the test cannot drift from the columns that ship."""
    spec = importlib.util.spec_from_file_location("zz_revision_0004_seed", MIGRATION)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    names = (
        "key", "label_de", "priority_rank", "excel_column", "isad",
        "dublin_core", "ead3", "value_type", "owning_agent", "description_de",
    )  # fmt: skip
    return [
        FieldDefinition(**dict(zip(names, row, strict=True))) for row in module.FIELD_DEFINITIONS
    ]


# --- fixture: one batch, three photos --------------------------------------------


class _Builder:
    def __init__(self, session: Any, org: Organization, fonds: Fonds, batch: UploadBatch) -> None:
        self.session = session
        self.org = org
        self.fonds = fonds
        self.batch = batch

    def asset(self, ref: str, sha: str, **extra: Any) -> Asset:
        asset = Asset(
            id=uuid.uuid4(),
            org_id=self.org.id,
            fonds_id=self.fonds.id,
            batch_id=self.batch.id,
            external_ref=ref,
            priority=1,
            **extra,
        )
        self.session.add(asset)
        self.session.add(
            AssetFile(
                org_id=self.org.id,
                asset_id=asset.id,
                side=AssetSide.RECTO,
                role=FileRole.MASTER,
                bucket="originals",
                s3_key=f"foto21/{ref}.jpg",
                mime_type="image/jpeg",
                bytes=4_200_000,
                sha256=sha,
            )
        )
        return asset

    def run(self, asset: Asset, public_id: str, outcome: str | None) -> Run:
        run = Run(
            id=uuid.uuid4(),
            org_id=self.org.id,
            public_id=public_id,
            asset_id=asset.id,
            batch_id=self.batch.id,
            status=RunStatus.AWAITING_REVIEW,
            outcome_confidence=Decimal(outcome) if outcome else None,
        )
        self.session.add(run)
        return run

    def evidence(
        self, run: Run, local_id: str, source_class: SourceClass, **extra: Any
    ) -> EvidenceItem:
        item = EvidenceItem(
            id=uuid.uuid4(),
            org_id=self.org.id,
            run_id=run.id,
            local_id=local_id,
            kind=EvidenceKind.EXTERNAL,
            source_class=source_class,
            agent_id=AgentId.GEO,
            observation=f"Beobachtung {local_id}",
            strength=Decimal("0.900"),
            **extra,
        )
        self.session.add(item)
        return item

    def version(
        self, record: CatalogueRecord, number: int, run: Run | None, user: User | None = None
    ) -> RecordVersion:
        version = RecordVersion(
            id=uuid.uuid4(),
            org_id=self.org.id,
            record_id=record.id,
            version=number,
            run_id=run.id if run else None,
            actor_kind=ActorKind.HUMAN if user else ActorKind.MACHINE,
            actor_user_id=user.id if user else None,
            change_kind=VersionChange.EDITED if user else VersionChange.GENERATED,
        )
        self.session.add(version)
        return version

    def field(
        self,
        version: RecordVersion,
        key: str,
        text: str | None,
        confidence: str | None,
        *,
        decision: FieldDecisionKind = FieldDecisionKind.EMIT,
        evidence: tuple[EvidenceItem, ...] = (),
        **extra: Any,
    ) -> RecordField:
        rf = RecordField(
            id=uuid.uuid4(),
            org_id=self.org.id,
            record_version_id=version.id,
            field_key=key,
            decision=decision,
            value_text=text,
            confidence=Decimal(confidence) if confidence else None,
            raw_confidence=Decimal(confidence) if confidence else None,
            corroborating_count=len(evidence),
            source_classes=sorted({str(e.source_class) for e in evidence}),
            generated_by="adjudicator",
            **extra,
        )
        self.session.add(rf)
        for item in evidence:
            self.session.add(
                RecordFieldEvidence(record_field_id=rf.id, evidence_id=item.id, org_id=self.org.id)
            )
        return rf


@pytest.fixture
async def batch(app) -> dict[str, Any]:
    database = app.state.database
    async with database.session(TenantContext.bootstrap()) as session:
        for definition in _field_definitions():
            session.add(definition)
        org = Organization(
            id=uuid.uuid4(), slug="stasz", name="Staatsarchiv Schwyz", kind=OrganizationKind.ARCHIVE
        )
        session.add(org)
        reviewer = User(
            id=uuid.uuid4(),
            org_id=org.id,
            cognito_sub="sub-rr",
            cognito_username="rr",
            email="r.ruch@sz.ch",
            full_name="Ralph Ruch",
        )
        session.add(reviewer)
        fonds = Fonds(
            id=uuid.uuid4(),
            org_id=org.id,
            slug="foto21",
            reference_code="SG.CIV",
            title="Fotografien 21. Jahrhundert",
            level=DescriptionLevel.SERIES,
            extent_unit=ExtentUnit.PHOTOGRAPHS,
        )
        session.add(fonds)
        upload = UploadBatch(
            id=uuid.uuid4(), org_id=org.id, source=BatchSource.ZIP, label="prio1"
        )
        session.add(upload)
        await session.flush()
        b = _Builder(session, org, fonds, upload)

        # Inserted out of order: the workbook must sort them naturally.
        c = b.asset("SG.CIV.17.20", "c" * 64)
        bb = b.asset("SG.CIV.17.11", "b" * 64)
        a = b.asset(
            "SG.CIV.17.3",
            "a" * 64,
            captured_at=datetime(2014, 6, 12, 10, 30, tzinfo=UTC),
        )
        await session.flush()

        # (a) finished machine run, one field at Prüfbedarf, one left open.
        run_a = b.run(a, "run_A17003", "0.810")
        geo = b.evidence(
            run_a, "ev_001", SourceClass.GEODATA_LAYER, source="swisstopo", source_uri=URL
        )
        ref = b.evidence(
            run_a, "ev_002", SourceClass.REFERENCE_IMAGE, source="Wikimedia Commons",
            source_uri="javascript:alert(1)", model_id="eu.anthropic.claude-opus-5-5",
        )  # fmt: skip
        web = b.evidence(run_a, "ev_003", SourceClass.OPEN_WEB, source="=HYPERLINK(\"x\")")
        record_a = CatalogueRecord(
            id=uuid.uuid4(), org_id=org.id, asset_id=a.id, fonds_id=fonds.id, current_version=1
        )
        session.add(record_a)
        await session.flush()
        v_a = b.version(record_a, 1, run_a)
        await session.flush()
        b.field(
            v_a, "locality", "Schwyz", "0.930", evidence=(geo, ref),
            explanation=(
                "2 Quellenarten bestätigen (Geodaten, Referenzbild) · 0 anerkannte Einwände"
            ),
        )  # fmt: skip
        b.field(
            v_a, "municipality", "Schwyz", "0.950", evidence=(geo,),
            value={"name": "Schwyz", "bfs_nr": 1372},
        )  # fmt: skip
        b.field(
            v_a, "building", "Alterszentrum Acherhof", "0.810", evidence=(ref, web),
            decision=FieldDecisionKind.EMIT_WITH_REVIEW,
            explanation="2 Quellenarten bestätigen (Referenzbild, Web) · 1 anerkannter Einwand",
        )  # fmt: skip
        b.field(
            v_a, "landscape_feature", None, None,
            decision=FieldDecisionKind.ABSTAIN,
            abstain_reason=AbstainReason.INSUFFICIENT_EVIDENCE,
        )  # fmt: skip
        b.field(
            v_a, "street", "Grundstrasse 32", "0.920", evidence=(geo,),
            explanation="1 Quellenart bestätigt (Geodaten) · 0 anerkannte Einwände",
        )  # fmt: skip
        b.field(
            v_a, "subject_point", "2692500, 1208300", "0.920", evidence=(geo,),
            value={"lv95": [2692500, 1208300]},
        )  # fmt: skip

        # (b) machine v1, then a human accepted every emitted field → v2.
        run_b = b.run(bb, "run_A17011", "0.910")
        geo_b = b.evidence(run_b, "ev_001", SourceClass.GEODATA_LAYER, source="swisstopo")
        record_b = CatalogueRecord(
            id=uuid.uuid4(), org_id=org.id, asset_id=bb.id, fonds_id=fonds.id, current_version=2
        )
        session.add(record_b)
        await session.flush()
        v1 = b.version(record_b, 1, run_b)
        v2 = b.version(record_b, 2, None, reviewer)
        await session.flush()
        for version in (v1, v2):
            human = (
                {"origin": FieldOrigin.HUMAN, "reviewed_by": reviewer.id,
                 "review_action": ReviewAction.ACCEPTED_UNCHANGED}
                if version is v2 else {}
            )  # fmt: skip
            b.field(version, "locality", "Brunnen", "0.910", evidence=(geo_b,), **human)
            b.field(version, "municipality", "Ingenbohl", "0.960", evidence=(geo_b,), **human)
            for key in ("building", "landscape_feature", "street"):
                b.field(
                    version, key, None, None,
                    decision=FieldDecisionKind.ABSTAIN,
                    abstain_reason=AbstainReason.INSUFFICIENT_EVIDENCE,
                )  # fmt: skip
        for key in ("locality", "municipality"):
            session.add(
                FieldReview(
                    org_id=org.id, record_id=record_b.id, from_version=1, to_version=2,
                    field_key=key, action=ReviewAction.ACCEPTED_UNCHANGED, user_id=reviewer.id,
                )
            )  # fmt: skip

        # (c) no run yet.
        _ = c
        batch_id = upload.id
    return {"batch_id": batch_id}


async def _workbook(app, batch_id: uuid.UUID) -> Any:
    async with app.state.database.session(TenantContext.bootstrap()) as session:
        data = await build_workbook(session, batch_id, generated_at=GENERATED_AT)
    return load_workbook(io.BytesIO(data))


def _rows(ws: Any) -> list[dict[str, Any]]:
    header = [cell.value for cell in ws[1]]
    rows = ws.iter_rows(min_row=2, values_only=True)
    return [dict(zip(header, row, strict=True)) for row in rows]


# --- the workbook ------------------------------------------------------------------


async def test_sheets_and_the_exact_header(app, batch) -> None:
    wb = await _workbook(app, batch["batch_id"])
    assert wb.sheetnames == ["Ergebnisse", "Belege", "Methode", "Spalten"]
    ws = wb["Ergebnisse"]
    assert [cell.value for cell in ws[1]] == HEADER
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref is not None
    assert ws["A1"].font.bold


async def test_one_row_per_photo_in_natural_order(app, batch) -> None:
    rows = _rows((await _workbook(app, batch["batch_id"]))["Ergebnisse"])
    assert [r["Dateiname"] for r in rows] == ["SG.CIV.17.3", "SG.CIV.17.11", "SG.CIV.17.20"]
    assert [r["Datei_SHA256"] for r in rows] == ["a" * 64, "b" * 64, "c" * 64]


async def test_the_machine_row(app, batch) -> None:
    a = _rows((await _workbook(app, batch["batch_id"]))["Ergebnisse"])[0]
    assert a["Ortschaft"] == "Schwyz"
    assert a["Ortschaft_Konfidenz"] == pytest.approx(0.93)
    assert a["Ortschaft_Quellen"].startswith("2 Quellenarten bestätigen (Geodaten, Referenzbild)")
    assert a["Gemeinde"] == "Schwyz"
    assert a["Gemeinde_BFS"] == 1372
    # No stored explanation: rendered from the source classes.
    assert a["Gemeinde_Quellen"] == "1 Quellenart bestätigt (Geodaten) · 0 anerkannte Einwände"
    assert a["Gebaeude"] == "Alterszentrum Acherhof"
    assert a["Gebaeude_Konfidenz"] == pytest.approx(0.81)
    # The abstained field: empty value, and an empty confidence cell, never 0.
    assert a["See_Fluss_Berg"] is None
    assert a["See_Fluss_Berg_Konfidenz"] is None
    assert a["Strasse"] == "Grundstrasse 32"
    assert (a["Motiv_LV95_E"], a["Motiv_LV95_N"]) == (2692500, 1208300)
    assert a["Aufnahmestandort_LV95_E"] is None
    assert a["Aufnahmerichtung_Grad"] is None
    assert a["Aufnahmedatum_EXIF"].date().isoformat() == "2014-06-12"
    assert a["Ausserhalb_SZ"] == "nein"
    assert a["Gesamtkonfidenz"] == pytest.approx(0.81)
    assert a["Status"] == "Prüfbedarf"
    assert a["Pruefbedarf"] == "Gebaeude"
    assert a["Unbestimmt"] == "See_Fluss_Berg (insufficient_evidence)"
    assert a["QC_Status"] == "maschinell"
    assert a["QC_Von"] is None
    assert a["Run_ID"] == "run_A17003"
    assert a["Erzeugt_am"] == "2026-10-24T09:02:44Z"


async def test_the_confirmed_row_and_the_empty_row(app, batch) -> None:
    wb = await _workbook(app, batch["batch_id"])
    ws = wb["Ergebnisse"]
    _, b, c = _rows(ws)
    assert b["Ortschaft"] == "Brunnen"
    assert b["Status"] == "übernahmefähig"
    assert b["Gesamtkonfidenz"] == pytest.approx(0.91)
    assert b["QC_Status"] == "geprüft"
    assert b["QC_Von"] == "Ralph Ruch"
    assert b["Run_ID"] == "run_A17011"  # the run behind the newest machine version
    assert b["Unbestimmt"] == (
        "Gebaeude (insufficient_evidence), See_Fluss_Berg (insufficient_evidence), "
        "Strasse (insufficient_evidence)"
    )
    filled = {k for k, v in c.items() if v is not None}
    assert filled == {"Dateiname", "Datei_SHA256", "Prio", "QC_Status"}
    assert c["QC_Status"] == "offen"
    status = HEADER.index("Status") + 1
    assert ws.cell(row=2, column=status).fill.fgColor.rgb.endswith("FFF2CC")
    assert ws.cell(row=3, column=status).fill.fgColor.rgb.endswith("E2F0D9")


async def test_belege_cites_only_the_head_version(app, batch) -> None:
    wb = await _workbook(app, batch["batch_id"])
    rows = _rows(wb["Belege"])
    # (a): locality 2 + municipality 1 + building 2 + street 1 + subject point 1;
    # (b): 2 from the head version v2, not the 2 behind v1; (c): none.
    assert len(rows) == 9
    assert [r["Feld"] for r in rows if r["Dateiname"] == "SG.CIV.17.3"][-1] == "Motiv_LV95"
    assert {r["Dateiname"] for r in rows} == {"SG.CIV.17.3", "SG.CIV.17.11"}
    first = rows[0]
    assert (first["Dateiname"], first["Feld"]) == ("SG.CIV.17.3", "Ortschaft")
    assert first["Quellenart"] in {"Geodaten", "Referenzbild"}
    ws = wb["Belege"]
    links = [ws.cell(row=r, column=5).hyperlink for r in range(2, ws.max_row + 1)]
    targets = {link.target for link in links if link is not None}
    assert targets == {URL}  # the javascript: URL is text, not a link
    # Untrusted text never becomes a formula.
    quelle = [ws.cell(row=r, column=4) for r in range(2, ws.max_row + 1)]
    injected = [cell for cell in quelle if str(cell.value).startswith("=")]
    assert injected and all(cell.data_type == "s" for cell in injected)


async def test_methode_states_the_rule_and_its_limits(app, batch) -> None:
    wb = await _workbook(app, batch["batch_id"])
    text = "\n".join(
        str(cell) for row in wb["Methode"].iter_rows(values_only=True) for cell in row if cell
    )
    assert "Konfidenz ≥ 0.90 = übernahmefähig" in text
    assert "unter 0.50 bleibt das Feld leer" in text
    assert "nicht kalibriert" in text
    assert "zehn bestätigende Quellen = 100 %" in text
    for column in HEADER:
        assert column in text
    assert "2026-10-24T09:02:44Z" in text


async def test_spalten_maps_every_field(app, batch) -> None:
    rows = _rows((await _workbook(app, batch["batch_id"]))["Spalten"])
    mapped = {r["Feld (field_key)"]: r for r in rows if r["Feld (field_key)"]}
    assert mapped["municipality"]["Excel-Spalte"] == "Gemeinde"
    assert mapped["date"]["ISAD(G)"] == "3.1.3"
    assert [r["Excel-Spalte"] for r in rows[:5]] == [
        "Ortschaft", "Gemeinde", "Gebaeude", "See_Fluss_Berg", "Strasse",
    ]  # fmt: skip


async def test_qc_summary(app, batch) -> None:
    async with app.state.database.session(TenantContext.bootstrap()) as session:
        summary = await qc_summary(session, batch["batch_id"])
    assert summary == {
        "batch": "prio1",
        "total": 3,
        "processed": 2,
        "confirmed": 1,
        "corrected": 0,
        "open": 1,
        "machine": 1,
        "overrideRate": {
            "locality": 0.0,
            "municipality": 0.0,
            "building": None,
            "landscape_feature": None,
            "street": None,
        },
    }


# --- QC_Status precedence (pure) ---------------------------------------------------


def _review(key: str, action: ReviewAction) -> FieldReview:
    return FieldReview(field_key=key, action=action)


EMITTED = ["locality", "municipality"]
OK = ReviewAction.ACCEPTED_UNCHANGED


def test_no_human_touch_is_maschinell() -> None:
    assert qc_status(run_finished=True, reviews=[], emitted_fields=EMITTED) == "maschinell"


def test_no_human_touch_and_nothing_emitted_is_still_maschinell() -> None:
    assert qc_status(run_finished=True, reviews=[], emitted_fields=[]) == "maschinell"


def test_every_emitted_field_accepted_is_geprueft() -> None:
    reviews = [_review("locality", OK), _review("municipality", OK)]
    assert qc_status(run_finished=True, reviews=reviews, emitted_fields=EMITTED) == "geprüft"


@pytest.mark.parametrize(
    "action", [ReviewAction.EDITED, ReviewAction.REJECTED, ReviewAction.FILLED]
)
def test_one_correction_beats_every_acceptance(action: ReviewAction) -> None:
    reviews = [_review("locality", OK), _review("municipality", OK), _review("street", action)]
    assert qc_status(run_finished=True, reviews=reviews, emitted_fields=EMITTED) == "korrigiert"


def test_a_correction_counts_even_while_a_rerun_is_in_progress() -> None:
    reviews = [_review("building", ReviewAction.EDITED)]
    assert qc_status(run_finished=False, reviews=reviews, emitted_fields=EMITTED) == "korrigiert"


def test_a_partial_review_is_offen() -> None:
    reviews = [_review("locality", OK)]
    assert qc_status(run_finished=True, reviews=reviews, emitted_fields=EMITTED) == "offen"


def test_an_unfinished_run_is_offen() -> None:
    assert qc_status(run_finished=False, reviews=[], emitted_fields=[]) == "offen"
